import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
from datetime import datetime, time, timedelta
import io

# ================= 核心逻辑函数 =================

def process_zhaji(file):
    """处理闸机记录"""
    df = pd.read_excel(file, sheet_name=0, dtype={"编号": str})
    df = df[["编号", "姓名", "部门", "日期", "考勤点"]].copy()
    df.rename(columns={"日期": "完整日期和时间"}, inplace=True)
    
    def extract_status(kaoqin_point):
        if pd.isna(kaoqin_point): return ""
        text = str(kaoqin_point).strip()
        length = len(text)
        if length == 16:
            status = text[5]
        elif length == 12:
            status = text[0]
        else:
            return ""
        return "进" if status == "入" else status

    df["进出状态"] = df["考勤点"].apply(extract_status)
    return df[["编号", "姓名", "部门", "完整日期和时间", "考勤点", "进出状态"]]

def process_banci(file):
    """处理排班表"""
    try:
        df = pd.read_excel(file, sheet_name=0, skiprows=1)
        if '班次' not in df.columns:
            file.seek(0)
            df = pd.read_excel(file, sheet_name=0)
    except:
        file.seek(0)
        df = pd.read_excel(file, sheet_name=0)
        
    df = df[["序号", "姓名", "日期", "班次"]].copy()
    df["班次"] = df["班次"].replace("未排班", "缺勤")
    return df

def calc_outside(df_zhaji, df_banci):
    """计算外出时长"""
    numeric_vals = pd.to_numeric(df_zhaji['完整日期和时间'], errors='coerce')
    dt_str = pd.to_datetime(df_zhaji['完整日期和时间'], errors='coerce')
    dt_num = pd.to_datetime(numeric_vals, unit='D', origin='1899-12-30', errors='coerce')
    df_zhaji['完整日期和时间'] = dt_str.fillna(dt_num)

    df_zhaji['日期'] = df_zhaji['完整日期和时间'].dt.strftime('%Y-%m-%d')
    df_zhaji['姓名'] = df_zhaji['姓名'].astype(str).str.strip()
    
    df_banci['姓名'] = df_banci['姓名'].astype(str).str.strip()
    df_banci['日期'] = df_banci['日期'].astype(str).str.strip()
    df_banci['班次'] = df_banci['班次'].astype(str).str.strip()

    df_banci.rename(columns={'班次': '班次_shift'}, inplace=True)
    merged_df = df_zhaji.merge(df_banci, on=['姓名', '日期'], how='left')
    merged_df['班次'] = merged_df['班次_shift'].fillna('缺勤')
    merged_df['班次'] = merged_df['班次'].replace('nan', '缺勤')

    valid_df = merged_df[merged_df['班次'].str.strip() == '白班'].copy()
    if len(valid_df) == 0: return pd.DataFrame()

    valid_df = valid_df.sort_values(by=['姓名', '完整日期和时间'])
    outside_records = []

    for (name, date), group in valid_df.groupby(['姓名', '日期']):
        sample_time = group['完整日期和时间'].iloc[0]
        win_start = datetime.combine(sample_time.date(), time(8, 0, 0))
        win_end = datetime.combine(sample_time.date(), time(20, 0, 0))
        
        last_out_time = None
        for _, row in group.iterrows():
            t = row['完整日期和时间']
            s = row['进出状态']
            
            if s == '出':
                if win_start <= t <= win_end:
                    last_out_time = t
            elif s == '进' and last_out_time is not None:
                duration_minutes = (t - last_out_time).total_seconds() / 60.0
                if duration_minutes > 2:
                    outside_records.append({
                        '姓名': name,
                        '日期': date,
                        '班次': '白班',
                        '外出开始时间': last_out_time,
                        '外出结束时间': t,
                        '外出时长(分钟)': round(duration_minutes, 1),
                        '外出时长(小时)': duration_minutes / 60.0 
                    })
                last_out_time = None
                
    if not outside_records: return pd.DataFrame()
    return pd.DataFrame(outside_records)

def generate_final_excel(result_df):
    """生成最终的Excel"""
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        result_df.to_excel(writer, sheet_name='源数据', index=False)
    output.seek(0)
    
    wb = openpyxl.load_workbook(output)
    ws_source = wb.worksheets[0]
    
    # 数据分析 1：按天汇总
    analysis_df = result_df.groupby(['姓名', '日期']).agg(
        外出次数=('外出时长(小时)', 'count'),
        外出时长总和=('外出时长(小时)', 'sum')
    ).reset_index()
    analysis_df.rename(columns={'外出时长总和': '外出时长总和(小时)'}, inplace=True)
    
    ws_analysis = wb.create_sheet('数据分析')
    ws_analysis.append(['姓名', '日期', '外出次数', '外出时长总和(小时)'])
    
    yellow_fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
    for index, row in analysis_df.iterrows():
        ws_analysis.append(list(row))
        if row['外出次数'] <= 2 and row['外出时长总和(小时)'] >= 8:
            current_row = ws_analysis.max_row
            for col_idx in range(1, 5):
                ws_analysis.cell(row=current_row, column=col_idx).fill = yellow_fill
        ws_analysis.cell(row=ws_analysis.max_row, column=4).number_format = '0.00'

    # 数据分析 2：按员工汇总
    employee_summary_df = result_df.groupby('姓名').agg(
        总外出次数=('外出时长(小时)', 'count'),
        总外出时长=('外出时长(小时)', 'sum')
    ).reset_index()
    employee_summary_df.rename(columns={'总外出时长': '总外出时长(小时)'}, inplace=True)
    
    start_col = 6
    title_cell = ws_analysis.cell(row=1, column=start_col, value="每位员工总外出汇总")
    title_cell.font = Font(name='微软雅黑', bold=True, size=11)
    
    headers_summary = ['姓名', '总外出次数', '总外出时长(小时)']
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    for col_offset, header in enumerate(headers_summary):
        cell = ws_analysis.cell(row=2, column=start_col + col_offset, value=header)
        cell.font = Font(name='微软雅黑', bold=True)
        cell.fill = header_fill

    for row_idx, row in employee_summary_df.iterrows():
        r = row_idx + 3
        ws_analysis.cell(row=r, column=start_col, value=row['姓名']).font = Font(name='微软雅黑')
        ws_analysis.cell(row=r, column=start_col + 1, value=row['总外出次数']).font = Font(name='微软雅黑')
        val_cell = ws_analysis.cell(row=r, column=start_col + 2, value=row['总外出时长(小时)'])
        val_cell.font = Font(name='微软雅黑')
        val_cell.number_format = '0.00'

    # 合计行
    total_count = employee_summary_df['总外出次数'].sum()
    total_hours = employee_summary_df['总外出时长(小时)'].sum()
    total_row_idx = len(employee_summary_df) + 3

    ws_analysis.cell(row=total_row_idx, column=start_col, value="合计").font = Font(name='微软雅黑', bold=True)
    ws_analysis.cell(row=total_row_idx, column=start_col + 1, value=total_count).font = Font(name='微软雅黑', bold=True)
    cell_total_hours = ws_analysis.cell(row=total_row_idx, column=start_col + 2, value=total_hours)
    cell_total_hours.font = Font(name='微软雅黑', bold=True)
    cell_total_hours.number_format = '0.00'

    # 全局格式化
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    for row in ws_analysis.iter_rows(min_row=2, max_row=len(employee_summary_df)+3, min_col=start_col, max_col=start_col+2):
        for cell in row:
            cell.border = thin_border

    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    current_size = cell.font.size if cell.font.size else 11
                    cell.font = Font(name='微软雅黑', size=current_size)
                    
    if '外出时长(小时)' in [c.value for c in ws_source[1]]:
        target_col_idx = [c.value for c in ws_source[1]].index('外出时长(小时)') + 1
        for row_idx in range(2, ws_source.max_row + 1):
            ws_source.cell(row=row_idx, column=target_col_idx).number_format = '0.00'

    final_output = io.BytesIO()
    wb.save(final_output)
    final_output.seek(0)
    return final_output

# ================= 页面展示 =================

st.set_page_config(page_title="外出时长一键统计", page_icon="📊")
st.title("📊 外出时长一键统计系统")
st.markdown("请上传 **闸机记录** 和 **排班表**，点击下方按钮即可自动生成最终结果。")

# --- 下载模板区域（已按最新截图精准还原格式和公式） ---
with st.expander("📥 点击下载上传模板（包含说明和标准示例，请直接填入数据后上传）"):
    col_t1, col_t2 = st.columns(2)
    
    # 动态生成闸机模板
    wb_zhaji = openpyxl.Workbook()
    ws_zhaji = wb_zhaji.active
    ws_zhaji.title = "Sheet1"
    
    zhaji_headers = ['编号', '姓名', '部门', '日期', '日期2', '时间', '考勤状态', '考勤点', '自定义名称', '数据来源', '处理类型', '体温', '体温异常']
    ws_zhaji.append(zhaji_headers)
    
    # 写入示例数据和复杂的 Excel 公式（填充至第 27 行）
    for row_idx in range(2, 28):
        if row_idx == 2:
            # 第一行包含示例数据
            ws_zhaji.cell(row=row_idx, column=1, value='001').number_format = '@' # 强制文本
            ws_zhaji.cell(row=row_idx, column=2, value='张**')
            ws_zhaji.cell(row=row_idx, column=3, value='制造')
            ws_zhaji.cell(row=row_idx, column=4, value=46150.461805555555)
            ws_zhaji.cell(row=row_idx, column=7, value='无')
            ws_zhaji.cell(row=row_idx, column=8, value='热充厂出口_门1_进门读卡器1')
            ws_zhaji.cell(row=row_idx, column=9, value='-')
            ws_zhaji.cell(row=row_idx, column=10, value='原始记录')
            ws_zhaji.cell(row=row_idx, column=11, value='-')
            ws_zhaji.cell(row=row_idx, column=12, value='-')
            ws_zhaji.cell(row=row_idx, column=13, value='-')
        
        # 循环写入 E 列和 F 列公式（自动匹配当前行号）
        formula_e = f'=IF(ISTEXT(D{row_idx}),DATEVALUE(SUBSTITUTE(LEFT(D{row_idx},FIND(" ",D{row_idx}&" ")-1),"-","/")),INT(D{row_idx}))'
        formula_f = f'=IFERROR(IF(ISTEXT(D{row_idx}),TIMEVALUE(TRIM(MID(D{row_idx},FIND(" ",D{row_idx})+1,99))),MOD(D{row_idx},1)),"")'
        ws_zhaji.cell(row=row_idx, column=5, value=formula_e)
        ws_zhaji.cell(row=row_idx, column=6, value=formula_f)
        
    template_zhaji = io.BytesIO()
    wb_zhaji.save(template_zhaji)
    template_zhaji.seek(0)
    
    # 动态生成排班模板
    wb_banci = openpyxl.Workbook()
    ws_banci = wb_banci.active
    ws_banci.title = "排班明细(1人1天1行)"
    
    # 大标题并合并单元格
    ws_banci.append(['梓阳俊-排班明细（1人1天1行,1天1班次1岗位,同名员工已合并）8:30左右=白班/20:30左右=晚班/空白=未排班'])
    ws_banci.merge_cells('A1:E1')
    
    # 表头
    ws_banci.append(['序号', '姓名', '日期', '班次', '岗位'])
    
    # 示例数据（根据最新截图修改）
    ws_banci.append([1, '刘*', '2026-05-03', '白班', '厂区/办公室清洁'])
    ws_banci.append([2, '华**', '2026-05-03', '白班', '厂区/办公室清洁'])
    ws_banci.append([3, '张**', '2026-05-03', '白班', '厂区/办公室清洁'])
    
    template_banci = io.BytesIO()
    wb_banci.save(template_banci)
    template_banci.seek(0)
    
    with col_t1:
        st.download_button(
            label="下载《闸机记录》模板",
            data=template_zhaji,
            file_name="闸机记录 - 模板.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    with col_t2:
        st.download_button(
            label="下载《排班表》模板",
            data=template_banci,
            file_name="排班表 - 模板.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    st.info("💡 提示：模板内已包含标准示例。闸机表的“日期2”和“时间”列已自动填充了适配至第 27 行的公式。")
# ----------------------------

col1, col2 = st.columns(2)
with col1:
    zhaji_file = st.file_uploader("1. 上传闸机记录 (Excel)", type=["xlsx"])
with col2:
    banci_file = st.file_uploader("2. 上传排班表 (Excel)", type=["xlsx"])

if st.button("🚀 一键开始处理", type="primary"):
    if not zhaji_file or not banci_file:
        st.warning("请先上传两个 Excel 文件！")
    else:
        with st.spinner("正在处理中，请稍候..."):
            try:
                df_zhaji = process_zhaji(zhaji_file)
                df_banci = process_banci(banci_file)
                result_df = calc_outside(df_zhaji, df_banci)
                
                if result_df.empty:
                    st.warning("没有找到任何符合条件的外出记录，请检查数据。")
                else:
                    final_excel = generate_final_excel(result_df)
                    st.success("✅ 处理完成！请点击下方按钮下载。")
                    st.download_button(
                        label="📥 下载 外出时长统计_白班版_最终修正.xlsx",
                        data=final_excel,
                        file_name="外出时长统计_白班版_最终修正.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
            except Exception as e:
                st.error(f"处理过程中出现错误：{e}")
